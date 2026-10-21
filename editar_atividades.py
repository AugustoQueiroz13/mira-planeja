# -*- coding: utf-8 -*-
"""
Ajusta atividades nas planilhas da pasta de entrada, sem edicao manual.

Uso:
  python editar_atividades.py dados/entrada            # so relatorio (nao altera)
  python editar_atividades.py dados/entrada --aplicar  # salva copias editadas

Em modo relatorio (padrao) nada e alterado: o script apenas mostra o que
faria. Com --aplicar, ele salva uma copia de cada planilha alterada com o
sufixo "_editado.xlsx", mantendo os originais intactos.
"""
import glob
import os
import sys
import unicodedata

import openpyxl


def norm(v):
    """minusculo, sem acento, sem espacos nas pontas"""
    s = str(v if v is not None else "")
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    return s.lower().strip()


# ------------------------------------------------------------- regras de edicao

# 1) Alterar a data realizada de atividades da 11.1 em Rio das Ostras
ALTERAR_DATA = [
    {"cod": "11.1", "local_contem": "cinamomo", "nova_data": "01/06/2026"},
    {"cod": "11.1", "local_contem": "esmeralda", "nova_data": "02/06/2026"},
]

# 2) Remover atividades 21.2 nao realizadas. Cada item casa quando TODAS as
#    palavras-chave aparecem na linha (atividade + municipio juntos).
REMOVER = [
    # Criterios ESPECIFICOS (seguros, casam uma atividade so):
    {"cod": "21.2", "chaves": ["conselho municipal de cultura", "cananeia"]},
    {"cod": "21.2", "chaves": ["sessao ordinaria", "carapebus"]},
    {"cod": "21.2", "chaves": ["sessao extraordinaria", "casimiro de abreu"]},
    {"cod": "21.2", "chaves": ["assembleia ordinaria", "orcamento participativo"]},
    # Criterios so por municipio (Ilha Comprida, Macae, Carapebus) foram
    # retirados porque ha varias 21.2 nesses municipios e removeriam demais.
    # Use o modo --listar 21.2 para ver todas e identificar as exatas, depois
    # adicione aqui criterios com palavras-chave da descricao de cada uma.
    # Ex: {"cod": "21.2", "chaves": ["ilha comprida", "<palavra da descricao>"]},
]


# ------------------------------------------------------------- deteccao de colunas

def acha_cabecalho(ws):
    """Procura a linha de cabecalho e mapeia as colunas pelo nome."""
    for r in range(1, 9):
        vals = {}
        for c in range(1, 40):
            cab = norm(ws.cell(r, c).value)
            if cab:
                vals[cab] = c
        # o cabecalho real tem "atividade" E alguma coluna de data (a linha de
        # titulo "Planilha de Atividades" tem so uma celula e e descartada)
        tem_atividade = any(k == "atividade" or k.startswith("atividade") for k in vals)
        tem_data = any("data" in k for k in vals)
        if tem_atividade and tem_data and len(vals) >= 4:
            def pega(*termos):
                for k, c in vals.items():
                    if any(t in k for t in termos):
                        return c
                return None
            cols = {
                "atividade": pega("atividade"),
                "data": pega("data realizada", "data da atividade", "data de realizacao"),
                "municipio": pega("municipio", "cidade"),
                "local": pega("local"),
            }
            if cols["atividade"]:
                return r, cols
    return None, None


def texto_linha(ws, r, cols):
    partes = []
    for campo in ("atividade", "municipio", "local"):
        c = cols.get(campo)
        if c:
            partes.append(norm(ws.cell(r, c).value))
    return " | ".join(partes)


def valor(ws, r, c):
    return str(ws.cell(r, c).value) if c and ws.cell(r, c).value is not None else ""


def cod_da_linha(ws, r, cols):
    return norm(ws.cell(r, cols["atividade"]).value)


# ------------------------------------------------------------- processamento

def processa(caminho, aplicar):
    wb = openpyxl.load_workbook(caminho)
    nome = os.path.basename(caminho)
    mudou = False
    relatorio = []

    for aba in wb.sheetnames:
        ws = wb[aba]
        hdr, cols = acha_cabecalho(ws)
        if not hdr:
            continue

        remover_linhas = []
        for r in range(hdr + 1, ws.max_row + 1):
            cod = cod_da_linha(ws, r, cols)
            if not cod:
                continue
            linha = texto_linha(ws, r, cols)

            # alterar data
            for regra in ALTERAR_DATA:
                if cod.startswith(regra["cod"]) and regra["local_contem"] in linha:
                    if cols["data"]:
                        antigo = ws.cell(r, cols["data"]).value
                        relatorio.append(
                            f"  [ALTERAR DATA] {aba} linha {r}: {antigo} -> {regra['nova_data']}"
                            f"  ({regra['local_contem']})"
                        )
                        if aplicar:
                            ws.cell(r, cols["data"]).value = regra["nova_data"]
                            mudou = True

            # remover
            for regra in REMOVER:
                if cod.startswith(regra["cod"]) and all(ch in linha for ch in regra["chaves"]):
                    atv = valor(ws, r, cols["atividade"])
                    mun = valor(ws, r, cols.get("municipio"))
                    dt = valor(ws, r, cols.get("data"))
                    relatorio.append(
                        f"  [REMOVER] {aba} L{r} | mun: {mun} | data: {dt[:10]}\n"
                        f"            \"{atv[:90]}\"\n"
                        f"            (casou regra: {regra['chaves']})"
                    )
                    remover_linhas.append(r)
                    break

        # remove de baixo para cima para nao baguncar os indices
        if aplicar:
            for r in sorted(set(remover_linhas), reverse=True):
                ws.delete_rows(r, 1)
                mudou = True

    if relatorio:
        print(f"\n== {nome} ==")
        for linha in relatorio:
            print(linha)

    if aplicar and mudou:
        saida = caminho.replace(".xlsx", "_editado.xlsx")
        wb.save(saida)
        print(f"  -> salvo: {os.path.basename(saida)}")

    return len(relatorio)


def lista_codigo(caminho, cod_alvo):
    """Lista todas as atividades de um codigo, para conferencia."""
    wb = openpyxl.load_workbook(caminho)
    nome = os.path.basename(caminho)
    achou = []
    for aba in wb.sheetnames:
        ws = wb[aba]
        hdr, cols = acha_cabecalho(ws)
        if not hdr:
            continue
        for r in range(hdr + 1, ws.max_row + 1):
            cod = cod_da_linha(ws, r, cols)
            if cod.startswith(cod_alvo):
                atv = valor(ws, r, cols["atividade"])
                mun = valor(ws, r, cols.get("municipio"))
                dt = valor(ws, r, cols.get("data"))
                achou.append(f"  {aba} L{r} | {dt[:10]} | {mun} | {atv[:85]}")
    if achou:
        print(f"\n== {nome} ==")
        for linha in achou:
            print(linha)
    return len(achou)


def main():
    pasta = sys.argv[1] if len(sys.argv) > 1 else "dados/entrada"
    aplicar = "--aplicar" in sys.argv

    # modo listagem: python editar_atividades.py dados/entrada --listar 21.2
    cod_listar = None
    if "--listar" in sys.argv:
        i = sys.argv.index("--listar")
        if i + 1 < len(sys.argv):
            cod_listar = norm(sys.argv[i + 1])
    if cod_listar:
        print(f"Listando todas as atividades do codigo {cod_listar} em {pasta}:")
        total = 0
        for caminho in sorted(glob.glob(os.path.join(pasta, "*.xlsx"))):
            if "_editado" not in caminho:
                total += lista_codigo(caminho, cod_listar)
        print(f"\nTotal encontrado: {total}")
        return

    modo = "APLICANDO (salva copias _editado)" if aplicar else "RELATORIO (nada e alterado)"
    print(f"Pasta: {pasta}  |  Modo: {modo}")

    total = 0
    for caminho in sorted(glob.glob(os.path.join(pasta, "*.xlsx"))):
        if "_editado" in caminho:
            continue
        total += processa(caminho, aplicar)

    print(f"\nTotal de ajustes encontrados: {total}")
    if not aplicar and total:
        print("Confira acima. Se estiver correto, rode de novo com --aplicar.")


if __name__ == "__main__":
    main()