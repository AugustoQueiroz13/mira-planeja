"""
Leitura das planilhas de comunicacao do PCS para o bloco 'comunicacao' do
dados.json:
 - Imprensa (clipping): materias publicadas, com veiculo, data e link.
 - Redes sociais: seguidores por rede ao longo dos meses.
 - Whatsapp: total de membros por mes.

A planilha de imprensa fica em dados/comunicacao/. Os seguidores e o Whatsapp
sao lidos da propria planilha de redes (abas especiais), encontrada na pasta
de entrada.
"""

import os
import re
import glob
import unicodedata

import pandas as pd
from openpyxl import load_workbook

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
_MESES_NORM = None


def _norm(c):
    f = unicodedata.normalize("NFKD", str(c or "").lower())
    return re.sub(r"\s+", " ", "".join(x for x in f if not unicodedata.combining(x))).strip()


def _meses_norm():
    global _MESES_NORM
    if _MESES_NORM is None:
        _MESES_NORM = {_norm(m) for m in MESES}
    return _MESES_NORM


def _txt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def _int(v):
    n = pd.to_numeric(v, errors="coerce")
    return int(n) if pd.notna(n) else 0


def _data_br(v):
    d = pd.to_datetime(v, dayfirst=True, errors="coerce")
    return d.strftime("%d/%m/%Y") if pd.notna(d) else _txt(v)


def _le_imprensa(caminho):
    materias = []
    abas = pd.read_excel(caminho, sheet_name=None, header=None)
    for _nome, raw in abas.items():
        hi = None
        for i in range(min(6, len(raw))):
            if any(_norm(x) == "tema do release" for x in raw.iloc[i].tolist()):
                hi = i
                break
        if hi is None:
            continue
        hdr = [_norm(c) for c in raw.iloc[hi].tolist()]

        def col(pref):
            for j, h in enumerate(hdr):
                if h.startswith(pref):
                    return j
            return None

        it, iv = col("tema do release"), col("veiculos que publicaram")
        itm, idv = col("tipo de midia"), col("data de veiculacao")
        imun, il = col("municipio do veiculo"), col("link da materia")
        for r in raw.iloc[hi + 1:].itertuples(index=False):
            r = list(r)

            def g(ix):
                return r[ix] if ix is not None and ix < len(r) else None

            tema = _txt(g(it))
            if not tema or tema.lower() in ("tema", "clipping", "tema do release"):
                continue
            materias.append({
                "tema": tema,
                "veiculo": _txt(g(iv)),
                "tipo_midia": _txt(g(itm)),
                "data": _data_br(g(idv)),
                "municipio": _txt(g(imun)),
                "link": _txt(g(il)),
            })
    return materias


def _le_seguidores(caminho):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    aba = next((s for s in wb.sheetnames if _norm(s).startswith("criacao e monitoramento")), None)
    if not aba:
        wb.close()
        return []
    rows = [list(r) for r in wb[aba].iter_rows(values_only=True)]
    wb.close()
    hi = next((i for i, r in enumerate(rows[:5]) if any(_norm(c) == "rede" for c in r)), None)
    if hi is None:
        return []
    hdr = [_norm(c) for c in rows[hi]]
    irede = hdr.index("rede")
    mesesmap = [(j, c.capitalize()) for j, c in enumerate(hdr) if c in _meses_norm()]
    org = ""
    out = []
    for r in rows[hi + 1:]:
        o = _txt(r[0]) if len(r) > 0 else ""
        if o:
            org = o
        rede = _txt(r[irede]) if irede < len(r) else ""
        if not rede or rede.lower() == "site":
            continue
        serie = [{"mes": mn, "valor": _int(r[j])} for j, mn in mesesmap if j < len(r) and _int(r[j]) > 0]
        if serie:
            out.append({"rede": rede, "org": org, "serie": serie})
    return out


def _le_whatsapp(caminho):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    aba = next((s for s in wb.sheetnames if _norm(s).startswith("whatsapp")), None)
    if not aba:
        wb.close()
        return {"por_mes": [], "total": 0}
    rows = [list(r) for r in wb[aba].iter_rows(values_only=True)]
    wb.close()
    hi = next((i for i, r in enumerate(rows[:5]) if any(_norm(c) == "municipio" for c in r)), None)
    if hi is None:
        return {"por_mes": [], "total": 0}
    hdr = [_norm(c) for c in rows[hi]]
    mesesmap = [(j, c.capitalize()) for j, c in enumerate(hdr) if c in _meses_norm()]
    somas = {mn: 0 for _, mn in mesesmap}
    for r in rows[hi + 1:]:
        for j, mn in mesesmap:
            if j < len(r):
                somas[mn] += _int(r[j])
    por_mes = [{"mes": mn, "membros": somas[mn]} for _, mn in mesesmap]
    total = max([p["membros"] for p in por_mes], default=0)
    return {"por_mes": por_mes, "total": total}


def _acha_redes(pasta):
    if not pasta or not os.path.isdir(pasta):
        return None
    for arq in sorted(glob.glob(os.path.join(pasta, "*.xlsx"))):
        if os.path.basename(arq).startswith("~$"):
            continue
        try:
            wb = load_workbook(arq, read_only=True)
            nomes = [_norm(s) for s in wb.sheetnames]
            wb.close()
            if any(n.startswith("criacao e monitoramento") for n in nomes):
                return arq
        except Exception:
            pass
    return None


def carrega_comunicacao(entrada, com_dir="dados/comunicacao"):
    materias = []
    if os.path.isdir(com_dir):
        for arq in sorted(glob.glob(os.path.join(com_dir, "*.xlsx"))):
            if os.path.basename(arq).startswith("~$"):
                continue
            try:
                materias += _le_imprensa(arq)
            except Exception:
                pass
    seguidores = []
    whatsapp = {"por_mes": [], "total": 0}
    redes = _acha_redes(entrada) or _acha_redes(com_dir)
    if redes:
        seguidores = _le_seguidores(redes)
        whatsapp = _le_whatsapp(redes)
    return {
        "imprensa": materias,
        "total_materias": len(materias),
        "seguidores": seguidores,
        "whatsapp": whatsapp,
    }